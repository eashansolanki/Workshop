"""Create genuine Office formats and convert via installed desktop applications."""
import json
import sys
import shutil
from pathlib import Path
from .core import path, process


def create_document(output: str, title: str, blocks: list[dict]) -> dict:
    """Create DOCX with paragraph, heading, bullet and table blocks."""
    from docx import Document
    target = path(output, new=True)
    if target.suffix.lower() != ".docx":
        raise ValueError("DOCX requires .docx; legacy .doc needs conversion")
    doc = Document()
    doc.add_heading(title, 0)
    for block in blocks:
        kind = block.get("type", "paragraph")
        if kind == "heading":
            doc.add_heading(str(block["text"]), min(4, max(1, int(block.get("level", 1)))))
        elif kind in {"paragraph", "bullet"}:
            doc.add_paragraph(str(block["text"]), style="List Bullet" if kind == "bullet" else None)
        elif kind == "table":
            rows = block["rows"]
            if not rows or not rows[0] or any(len(row) != len(rows[0]) for row in rows):
                raise ValueError("Tables require nonempty rectangular rows")
            table = doc.add_table(rows=0, cols=len(rows[0]))
            table.style = "Light Shading Accent 1"
            for row in rows:
                cells = table.add_row().cells
                for cell, value in zip(cells, row):
                    cell.text = str(value)
        else:
            raise ValueError(f"Unsupported block: {kind}")
    doc.save(target)
    return {"path": str(target), "format": "docx"}


def create_workbook(output: str, sheets: list[dict]) -> dict:
    """Create XLSX sheets with rows, explicit formulas, freeze panes and header styling."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.workbook.properties import CalcProperties
    target = path(output, new=True)
    if target.suffix.lower() != ".xlsx" or not sheets:
        raise ValueError("Provide .xlsx output and at least one sheet")
    names = [sheet["name"] for sheet in sheets]
    if len(set(name.casefold() for name in names)) != len(names):
        raise ValueError("Sheet names must be unique")
    book = Workbook()
    book.remove(book.active)
    book.calculation = CalcProperties(fullCalcOnLoad=True)
    for sheet in sheets:
        ws = book.create_sheet(sheet["name"])
        for row in sheet.get("rows", []):
            # Ordinary strings are literal. Formulas must be passed explicitly below.
            ws.append(row)
            for cell in ws[ws.max_row]:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    cell.data_type = "s"
        for address, formula in sheet.get("formulas", {}).items():
            if not isinstance(formula, str) or not formula.startswith("="):
                raise ValueError("Explicit formulas must start with =")
            if "[" in formula or "WEBSERVICE(" in formula.upper():
                raise ValueError("External workbook/network formulas are unsupported")
            ws[address] = formula
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="244A73")
        ws.freeze_panes = "A2"
        for column in ws.columns:
            ws.column_dimensions[column[0].column_letter].width = min(60, max(12, max(len(str(c.value or "")) for c in column) + 2))
    book.save(target)
    return {"path": str(target), "format": "xlsx", "formula_note": "Excel/Numbers must recalculate formulas; no cached results supplied"}


def create_slides(output: str, title: str, slides: list[dict]) -> dict:
    """Create editable PPTX title/content slides and speaker notes."""
    from pptx import Presentation
    from pptx.util import Inches
    target = path(output, new=True)
    if target.suffix.lower() != ".pptx":
        raise ValueError("PowerPoint output requires .pptx")
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
    front = deck.slides.add_slide(deck.slide_layouts[0])
    front.shapes.title.text = title
    for item in slides:
        slide = deck.slides.add_slide(deck.slide_layouts[1])
        slide.shapes.title.text = item["title"]
        text = slide.placeholders[1].text_frame
        text.clear()
        for i, bullet in enumerate(item.get("bullets", [])):
            paragraph = text.paragraphs[0] if i == 0 else text.add_paragraph()
            paragraph.text = str(bullet)
        if item.get("notes"):
            slide.notes_slide.notes_text_frame.text = str(item["notes"])
    deck.save(target)
    return {"path": str(target), "format": "pptx"}


def inspect_office(source: str) -> dict:
    """Extract bounded Office content to verify structure and enable edits via code tools."""
    file = path(source)
    suffix = file.suffix.lower()
    if suffix == ".docx":
        from docx import Document
        doc = Document(file)
        return {"paragraphs": [p.text for p in doc.paragraphs][:500],
                "tables": [[[c.text for c in r.cells] for r in t.rows] for t in doc.tables][:20]}
    if suffix == ".xlsx":
        from openpyxl import load_workbook
        book = load_workbook(file, read_only=True, data_only=False)
        try:
            return {"sheets": {ws.title: list(ws.iter_rows(max_row=100, max_col=50, values_only=True)) for ws in book}}
        finally:
            book.close()
    if suffix == ".pptx":
        from pptx import Presentation
        return {"slides": [[shape.text for shape in slide.shapes if shape.has_text_frame] for slide in Presentation(file).slides][:100]}
    raise ValueError("Inspection supports docx, xlsx, pptx")


def apple_convert(source: str, output: str) -> dict:
    """Import Office into native iWork or export iWork to Office/PDF using macOS."""
    if sys.platform != "darwin":
        raise RuntimeError("Pages/Numbers/Keynote require macOS")
    src, dst = path(source), path(output, new=True)
    mapping = {".docx": "Pages", ".pages": "Pages", ".xlsx": "Numbers", ".numbers": "Numbers", ".pptx": "Keynote", ".key": "Keynote"}
    native = {"Pages": ".pages", "Numbers": ".numbers", "Keynote": ".key"}
    office = {"Pages": (".docx", "Microsoft Word"), "Numbers": (".xlsx", "Microsoft Excel"), "Keynote": (".pptx", "Microsoft PowerPoint")}
    app = mapping.get(src.suffix.lower())
    if not app or dst.suffix.lower() not in {native[app], office[app][0], ".pdf"}:
        raise ValueError("Unsupported source/output pairing")
    # JSON string quoting matches AppleScript path string literals on macOS.
    src_literal, dst_literal = json.dumps(str(src), ensure_ascii=False), json.dumps(str(dst), ensure_ascii=False)
    if dst.suffix.lower() == native[app]:
        operation = f"save d in (POSIX file {dst_literal})"
    else:
        format_name = "PDF" if dst.suffix.lower() == ".pdf" else office[app][1]
        operation = f"export d to (POSIX file {dst_literal}) as {format_name}"
    script = f'''tell application "{app}"
set d to open (POSIX file {src_literal})
try
{operation}
close d saving no
on error messageText number errorNumber
try
close d saving no
end try
error messageText number errorNumber
end try
end tell'''
    result = process(["osascript", "-e", script], 120)
    if result["exit_code"] or not dst.exists():
        raise RuntimeError(result["stderr"] or "iWork did not produce the output")
    return {"path": str(dst), "format": dst.suffix[1:]}


def libreoffice_convert(source: str, output: str) -> dict:
    """Use LibreOffice for genuine legacy DOC/PDF conversions; never rename extensions."""
    import tempfile
    src, dst = path(source), path(output, new=True)
    binary = shutil.which("soffice")
    if not binary and Path("/Applications/LibreOffice.app/Contents/MacOS/soffice").exists():
        binary = "/Applications/LibreOffice.app/Contents/MacOS/soffice"
    if not binary:
        raise RuntimeError("Install LibreOffice for legacy .doc and PDF conversion")
    if src.suffix.lower() not in {".docx", ".xlsx", ".pptx", ".doc", ".odt"} or dst.suffix.lower() not in {".doc", ".pdf", ".docx"}:
        raise ValueError("Unsupported conversion")
    if dst.suffix.lower() == ".doc" and src.suffix.lower() not in {".docx", ".odt"}:
        raise ValueError("DOC conversion requires a word-processing document")
    with tempfile.TemporaryDirectory(prefix="workshop-lo-") as temporary:
        temp = Path(temporary)
        profile = (temp / "profile").as_uri()
        result = process([binary, f"-env:UserInstallation={profile}", "--headless", "--convert-to", dst.suffix[1:], "--outdir", str(temp), str(src)], 120)
        converted = temp / (src.stem + dst.suffix)
        if result["exit_code"] or not converted.is_file():
            raise RuntimeError(result["stderr"] or result["stdout"] or "Conversion failed")
        shutil.copyfile(converted, dst)
    return {"path": str(dst)}
